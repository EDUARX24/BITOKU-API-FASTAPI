import logging
from io import BytesIO

import pandas as pd
from fastapi import APIRouter, Depends, Query, Security, status
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db
from app.schemas.ApiResponse.response import ApiResponse
from app.schemas.paciente.paciente import PacienteCreate, PacienteUpdate
from app.services.paciente_service import (
    create_paciente_service,
    delete_paciente_service,
    export_pacientes_rows_service,
    get_paciente_service,
    list_pacientes_service,
    update_paciente_service,
)


logger = logging.getLogger(__name__)
router = APIRouter()


@router.post("/paciente", response_model=ApiResponse, status_code=status.HTTP_201_CREATED, tags=["Paciente"])
def paciente(
    paciente_in: PacienteCreate,
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    paciente_out = create_paciente_service(paciente_in, db, current_user)
    return ApiResponse(
        statusCode=201,
        isSuccess=True,
        errorMessages=[],
        successMessages=["Perfil de paciente creado exitosamente"],
        result=paciente_out,
    )


@router.get("/pacientes", response_model=ApiResponse, status_code=status.HTTP_200_OK, tags=["Paciente"])
def pacientes(
    pageSize: int = Query(0, ge=0),
    pageNumber: int = Query(1, ge=1),
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    pacientes_out = list_pacientes_service(pageSize, pageNumber, db, current_user)
    return ApiResponse(
        statusCode=200,
        isSuccess=True,
        errorMessages=[],
        successMessages=[],
        result=pacientes_out,
    )


@router.get("/paciente/{idPaciente}", response_model=ApiResponse, status_code=status.HTTP_200_OK, tags=["Paciente"])
def paciente_id(
    idPaciente: int,
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    paciente_out = get_paciente_service(idPaciente, db, current_user)
    return ApiResponse(
        statusCode=200,
        isSuccess=True,
        errorMessages=[],
        successMessages=[],
        result=paciente_out,
    )


@router.delete("/paciente/{idPaciente}", response_model=ApiResponse, status_code=status.HTTP_200_OK, tags=["Paciente"])
def eliminar_paciente(
    idPaciente: int,
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    paciente_out = delete_paciente_service(idPaciente, db, current_user)
    return ApiResponse(
        statusCode=200,
        isSuccess=True,
        errorMessages=[],
        successMessages=["Perfil de paciente eliminado correctamente"],
        result=paciente_out,
    )


@router.put("/paciente/{idPaciente}", response_model=ApiResponse, status_code=status.HTTP_200_OK, tags=["Paciente"])
def actualizar_paciente(
    idPaciente: int,
    paciente_in: PacienteCreate,
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    paciente_out = update_paciente_service(
        idPaciente,
        paciente_in,
        db,
        current_user,
        partial=False,
    )
    return ApiResponse(
        statusCode=200,
        isSuccess=True,
        errorMessages=[],
        successMessages=["Perfil de paciente actualizado correctamente"],
        result=paciente_out,
    )


@router.patch("/paciente/{idPaciente}", response_model=ApiResponse, status_code=status.HTTP_200_OK, tags=["Paciente"])
def actualizar_paciente_patch(
    idPaciente: int,
    paciente_in: PacienteUpdate,
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    paciente_out = update_paciente_service(
        idPaciente,
        paciente_in,
        db,
        current_user,
        partial=True,
    )
    return ApiResponse(
        statusCode=200,
        isSuccess=True,
        errorMessages=[],
        successMessages=["Perfil de paciente actualizado correctamente"],
        result=paciente_out,
    )


@router.get("/pacientes/export", status_code=status.HTTP_200_OK, tags=["Paciente"])
def exportar_pacientes_excel(
    pageSize: int = Query(0, ge=0),
    pageNumber: int = Query(1, ge=1),
    db: Session = Depends(get_db),
    current_user: dict = Security(get_current_user),
):
    data = export_pacientes_rows_service(pageSize, pageNumber, db, current_user)
    df = pd.DataFrame(data)

    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Pacientes", startrow=4)

        ws = writer.sheets["Pacientes"]
        ws.merge_cells("A1:E1")
        ws["A1"] = "CLÍNICA BITOKU"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:E2")
        ws["A2"] = "Listado de Pacientes"
        ws["A2"].font = Font(bold=True, size=12)
        ws["A2"].alignment = Alignment(horizontal="center")

        header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        header_font = Font(bold=True)
        for col_num, cell in enumerate(ws[5], 1):
            cell.fill = header_fill
            cell.font = header_font
            col_letter = get_column_letter(col_num)
            max_length = max(
                len(str(cell.value)) if cell.value else 0,
                *(len(str(row.value)) if row.value else 0 for row in ws[col_letter]),
            )
            ws.column_dimensions[col_letter].width = max_length + 2

    stream.seek(0)
    logger.info("Pacientes exportados por %s (%s registros)", current_user["idUsuario"], len(data))
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=pacientes.xlsx"},
    )
